"""Extract bounded evidence blocks with locations, then create parent/child chunks."""
import csv, io, json, re
from email import policy
from email.parser import BytesParser
from pathlib import Path
from uuid import uuid4
from lxml import html

SUPPORTED = {'.txt','.md','.eml','.pdf','.docx','.csv','.xlsx','.json'}
class ParseError(ValueError): pass

def parse(path: Path):
    ext=path.suffix.lower(); blocks=[]
    def add(text, locator, section='Document'):
        if text.strip(): blocks.append({'text':text.strip(),'locator':locator,'section':section})
    if ext in {'.txt','.md'}:
        section='Document'
        for i,part in enumerate(re.split(r'\n\s*\n',path.read_text(encoding='utf-8-sig')),1):
            if re.match(r'^#{1,6}\s',part):section=part.splitlines()[0].lstrip('# ')
            add(part, f'paragraph {i}',section)
    elif ext=='.eml':
        msg=BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        add('\n'.join(f'{k}: {msg.get(k, "")}' for k in ['From','To','Date','Subject']),'email headers','Headers')
        body=msg.get_body(preferencelist=('plain','html'))
        if body:
            content=body.get_content()
            if body.get_content_type()=='text/html':
                tree=html.fromstring(content)
                for el in tree.xpath('//script|//style'):el.drop_tree()
                content=tree.text_content()
            for i,p in enumerate(re.split(r'\n\s*\n',content),1):add(p,f'email body paragraph {i}','Email body')
    elif ext=='.pdf':
        from pypdf import PdfReader
        for n,page in enumerate(PdfReader(path).pages,1):add(page.extract_text() or '',f'page {n}',f'Page {n}')
        if not blocks:raise ParseError('needs_ocr: this PDF has no extractable text')
    elif ext=='.docx':
        from docx import Document
        doc=Document(path);section='Document'
        for i,p in enumerate(doc.paragraphs,1):
            if p.style.name.startswith('Heading'):section=p.text
            add(p.text,f'paragraph {i}',section)
        for n,table in enumerate(doc.tables,1):
            header=' | '.join(c.text for c in table.rows[0].cells)
            for i,row in enumerate(table.rows[1:],2):add(header+'\n'+' | '.join(c.text for c in row.cells),f'table {n} row {i}',f'Table {n}')
    elif ext=='.csv':
        reader=csv.DictReader(io.StringIO(path.read_text(encoding='utf-8-sig')))
        for i,row in enumerate(reader,2):add(' | '.join(f'{k}: {v}' for k,v in row.items()),f'row {i}','Table')
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(path,read_only=True,data_only=False)
        for ws in wb:
            rows=ws.iter_rows(values_only=True);header=next(rows,())
            for i,row in enumerate(rows,2):
                add(' | '.join(f'{k}: {v}' for k,v in zip(header,row) if v is not None),f'{ws.title} row {i}',ws.title)
        wb.close()
    elif ext=='.json':
        add(json.dumps(json.loads(path.read_text()),ensure_ascii=False,indent=2),'JSON document')
    else:raise ParseError(f'Unsupported extension: {ext}')
    if not blocks:raise ParseError('No text found')
    return blocks

# Chinese characters and English words form length units; paragraph/row boundaries take precedence.
UNITS=re.compile(r'[\u3400-\u9fff]|[A-Za-z0-9_]+|[^\s]',re.UNICODE)
def windows(text, size=300, overlap=40):
    units=list(UNITS.finditer(text))
    for start in range(0,len(units),size-overlap):
        stop=min(start+size,len(units))
        yield text[units[start].start():units[stop-1].end()]
        if stop==len(units):break

def chunk(blocks, split_child=None):
    """One evidence block is a parent; long blocks get bounded parent windows."""
    parents=[]; children=[]
    for b in blocks:
        for parent_text in windows(b['text'],size=1400,overlap=80):
            parent_id=str(uuid4())
            parents.append({'id':parent_id,'text':parent_text,'section':b['section'],'locator':b['locator']})
            for text in (split_child(parent_text) if split_child else windows(parent_text)):
                children.append({'id':str(uuid4()),'parent_id':parent_id,'text':text,'section':b['section'],'locator':b['locator']})
    return parents,children

def lexical_tokens(text):
    parts=re.findall(r'[a-z0-9]+|[\u3400-\u9fff]+',text.lower()); tokens=[]
    for p in parts:
        if re.match(r'[\u3400-\u9fff]',p):tokens.extend([p[i:i+2] for i in range(len(p)-1)] or [p])
        else:tokens.append(p)
    return list(dict.fromkeys(tokens))
