from tradex_rag.config import Settings
from tradex_rag.service import Library
from test_library import TestEmbedding
import time

def test_restart_preserves_vectors_and_recovers_queue(tmp_path):
    settings=Settings(data_dir=tmp_path)
    first=Library(settings,embedding=TestEmbedding())
    doc=first.create('terms.txt','已确认的付款条款'.encode(),'条款','a')
    first.process_next()
    queued=first.create('later.txt','下一笔付款'.encode(),'后续','a')
    with first.db.connect(write=True) as db:db.execute("UPDATE jobs SET status='running' WHERE id=?",(queued['job_id'],))
    first.close()
    second=Library(settings,embedding=TestEmbedding())
    try:
        assert second.search('付款','a',mode='vector')['results'][0]['document_id']==doc['document_id']
        second.start_worker()
        for _ in range(100):
            if second.document(queued['document_id'],'a')['current_version']==1:break
            time.sleep(.02)
        assert second.document(queued['document_id'],'a')['current_version']==1
    finally:second.close()

def test_scanned_pdf_has_explicit_parse_status(tmp_path):
    from pypdf import PdfWriter
    from tradex_rag.parsing import parse,ParseError
    import pytest
    p=tmp_path/'scan.pdf';writer=PdfWriter();writer.add_blank_page(width=600,height=800)
    with p.open('wb') as f:writer.write(f)
    with pytest.raises(ParseError,match='needs_ocr'):parse(p)

def test_docx_and_xlsx_rows_keep_headers(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from tradex_rag.parsing import parse
    d=Document();d.add_heading('Payment',level=1);d.add_paragraph('Buyer promise is not receipt.')
    table=d.add_table(rows=2,cols=2);table.cell(0,0).text='Invoice';table.cell(0,1).text='Amount';table.cell(1,0).text='INV-1';table.cell(1,1).text='6000'
    p=tmp_path/'terms.docx';d.save(p)
    assert any('Invoice | Amount' in b['text'] and '6000' in b['text'] for b in parse(p))
    wb=Workbook();ws=wb.active;ws.append(['Invoice','Amount']);ws.append(['INV-1','=1+2'])
    p=tmp_path/'ledger.xlsx';wb.save(p)
    assert 'Amount: =1+2' in parse(p)[0]['text']

def test_bge_child_windows_fit_actual_model_tokenizer():
    from tradex_rag.embedding import LocalBGE
    from tradex_rag.config import Settings
    import pytest
    settings=Settings()
    if not settings.model_dir.exists():pytest.skip('Download the optional local model to run this integration check')
    model=LocalBGE(settings.model_dir)
    text='供应商付款条款需要确认。 Payment responsibility requires documentary evidence. '*200
    pieces=list(model.split_text(text))
    assert len(pieces)>1
    assert all(len(model._tokenizer.encode(p,add_special_tokens=False).ids)<=350 for p in pieces)
